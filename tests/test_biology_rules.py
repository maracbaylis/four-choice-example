from __future__ import annotations

from pathlib import Path
import sys
import unittest

import pandas as pd
from openpyxl import Workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "scripts"))

from four_choice_core import (  # noqa: E402
    add_sequence_metrics,
    extract_metadata,
    metadata_qc,
    order_animal_output_columns,
    pca_table,
    parse_entry_sequence,
    reversal_error_type,
    trials_to_criterion,
)


class BiologyRuleTests(unittest.TestCase):
    def test_entry_sequence_digits_map_to_quadrants_and_odors(self) -> None:
        quadrants, odors = parse_entry_sequence("1234", "ABCD")

        self.assertEqual(quadrants, ["NW", "NE", "SW", "SE"])
        self.assertEqual(odors, ["A", "B", "C", "D"])

    def test_entry_sequence_uses_trial_specific_bowl_order(self) -> None:
        quadrants, odors = parse_entry_sequence("1234", "BCDA")

        self.assertEqual(quadrants, ["NW", "NE", "SW", "SE"])
        self.assertEqual(odors, ["B", "C", "D", "A"])

    def test_ttc_is_8_correct_out_of_10_non_omission_trials(self) -> None:
        trials = pd.DataFrame(
            {
                "animal_id": ["animal_1"] * 11,
                "phase": ["discrimination"] * 11,
                "trial": list(range(1, 12)),
                "correct": [True, True, False, True, False, True, True, False, True, True, True],
                "omission": [False, False, False, False, True, False, False, False, False, False, False],
            }
        )

        scored = add_sequence_metrics(trials)

        self.assertEqual(scored.loc[scored["trial"] == 11, "correct_in_last_10_non_omission"].item(), 8)
        self.assertEqual(trials_to_criterion(scored), 11)

    def test_reversal_a_errors_switch_from_perseverative_to_regressive(self) -> None:
        self.assertEqual(reversal_error_type("A", previous_rewarded="A", new_rewarded="B", seen_correct=False), "Perseverative")
        self.assertEqual(reversal_error_type("A", previous_rewarded="A", new_rewarded="B", seen_correct=True), "Regressive")
        self.assertEqual(reversal_error_type("C", previous_rewarded="A", new_rewarded="B", seen_correct=True), "Irrelevant")
        self.assertEqual(reversal_error_type("D", previous_rewarded="A", new_rewarded="B", seen_correct=True), "Novel")

    def test_metadata_qc_flags_missing_required_fields(self) -> None:
        metadata = pd.DataFrame(
            [
                {
                    "animal_id": "animal_1",
                    "experiment_cohort": "ARC",
                    "adversity_condition": "N/A",
                    "sex": "",
                    "genotype": "Wildtype",
                    "lab_source": "Example Lab",
                    "animal_id_inferred_from_filename": False,
                }
            ]
        )

        qc = metadata_qc(metadata)

        self.assertIn("missing_required_metadata", set(qc["qc_flag"]))
        self.assertIn("sex", set(qc["field"]))

    def test_animal_outputs_start_with_metadata_grouping_columns(self) -> None:
        animals = pd.DataFrame(
            [
                {
                    "rev_total_errors": 2,
                    "TTC_in_reversal": 18,
                    "source_file": "example.xlsx",
                    "lab_source": "Example Lab",
                    "animal_id": "animal_1",
                    "strain": "C57BL/6J",
                    "dob": "2026-01-01",
                    "homecage_mice_at_testing": 3,
                    "adversity_condition": "Control",
                    "genotype": "Wildtype",
                    "birth_litter_size_p05_20": 6,
                    "homecage_rank": 1,
                    "adversity_age_range": "N/A",
                    "vendor": "Example Vendor",
                    "unmapped_noise_variable": "tracked",
                }
            ]
        )

        ordered = order_animal_output_columns(animals)

        self.assertEqual(
            list(ordered.columns[:11]),
            [
                "animal_id",
                "source_file",
                "strain",
                "dob",
                "homecage_mice_at_testing",
                "adversity_condition",
                "genotype",
                "birth_litter_size_p05_20",
                "homecage_rank",
                "adversity_age_range",
                "vendor",
            ],
        )
        self.assertLess(ordered.columns.get_loc("unmapped_noise_variable"), ordered.columns.get_loc("TTC_in_reversal"))
        self.assertLess(ordered.columns.get_loc("TTC_in_reversal"), ordered.columns.get_loc("rev_total_errors"))

    def test_metadata_parser_keeps_unmapped_fields(self) -> None:
        workbook = Workbook()
        ws = workbook.active
        ws.title = "METADATA"
        ws["A1"] = "animal_id"
        ws["B1"] = "animal_1"
        ws["A2"] = "new possible noise variable"
        ws["B2"] = "keep me"

        metadata = extract_metadata(workbook)

        self.assertEqual(metadata["animal_id"], "animal_1")
        self.assertEqual(metadata["new_possible_noise_variable"], "keep me")

    def test_pca_scores_keep_metadata_context(self) -> None:
        animals = pd.DataFrame(
            [
                {
                    "animal_id": "animal_1",
                    "source_file": "one.xlsx",
                    "strain": "C57BL/6J",
                    "adversity_condition": "Control",
                    "genotype": "Wildtype",
                    "TTC_in_discrimination": 10,
                    "Total_errors_in_discrimination": 2,
                    "TTC_in_reversal": 20,
                    "Total_errors_in_reversal": 8,
                },
                {
                    "animal_id": "animal_2",
                    "source_file": "two.xlsx",
                    "strain": "C57BL/6J",
                    "adversity_condition": "Treatment",
                    "genotype": "Wildtype",
                    "TTC_in_discrimination": 14,
                    "Total_errors_in_discrimination": 4,
                    "TTC_in_reversal": 28,
                    "Total_errors_in_reversal": 12,
                },
            ]
        )

        scores, _, _ = pca_table(animals)

        self.assertEqual(list(scores.columns[:5]), ["animal_id", "source_file", "strain", "adversity_condition", "genotype"])
        self.assertIn("PC1", scores.columns)


if __name__ == "__main__":
    unittest.main()
