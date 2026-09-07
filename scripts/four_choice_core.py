#!/usr/bin/env python3
"""Parsing and summary code for ARC Four Choice TABS workbooks.

All derived fields are calculated in Python from user-entered workbook values.
The workbook SUMMARY sheet and helper formulas are not used for analysis.
"""

from __future__ import annotations

import datetime as dt
import math
import re
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

try:
    from openpyxl.worksheet.formula import ArrayFormula
except ImportError:
    ArrayFormula = ()


MISSING_TEXT = {"", "N/A", "NA", "NAN", "NONE", "-"}
NO_VALUE_TEXT = {*MISSING_TEXT, "UNKNOWN"}

TRIAL_SHEETS = {
    "DISC": "discrimination",
    "RECALL": "recall",
    "REVERSAL": "reversal",
    "REV": "reversal",
}

ODOR_TO_QUADRANT = {
    "O1, O2, O3, O4": {"O1": "NW", "O2": "NE", "O3": "SW", "O4": "SE"},
    "O2, O3, O4, O1": {"O1": "SE", "O2": "NW", "O3": "NE", "O4": "SW"},
    "O3, O4, O1, O2": {"O1": "SW", "O2": "SE", "O3": "NW", "O4": "NE"},
    "O4, O1, O2, O3": {"O1": "NE", "O2": "SW", "O3": "SE", "O4": "NW"},
    "ABCD": {"A": "NW", "B": "NE", "C": "SW", "D": "SE"},
    "BCDA": {"A": "SE", "B": "NW", "C": "NE", "D": "SW"},
    "CDAB": {"A": "SW", "B": "SE", "C": "NW", "D": "NE"},
    "DABC": {"A": "NE", "B": "SW", "C": "SE", "D": "NW"},
}

QUADRANT_NUMBER = {"1": "NW", "2": "NE", "3": "SW", "4": "SE"}
WEIGHT_STAGE_TO_PHASE = {
    "initial weighing": "baseline",
    "food restriction + handling": "handling",
    "food restriction + acclimation": "acclimation",
    "food restriction + shaping": "shaping",
    "food restriction + testing": "testing",
    "testing": "testing",
}

METADATA_ALIASES = {
    "animal_id": ["animal_id_xxx001", "animal_xxx001", "animal_id", "animal"],
    "strain": ["strain"],
    "dob": ["dob_mm_dd_yyyy", "dob"],
    "homecage_mice_at_testing": ["of_mice_in_homecage_at_testing"],
    "adversity_condition": ["adversity_condition"],
    "genotype": ["genotype"],
    "birth_litter_size_p05_20": ["of_mice_in_birth_litter_p05_20", "of_mice_in_birth_litter_p5_20"],
    "homecage_rank": ["rank_in_homecage"],
    "adversity_age_range": ["adversity_age_range_pxx_xx"],
    "vendor": ["vendor"],
    "p20_22_weight_g": ["p20_22_weight"],
    "homecage_nesting_material": ["home_cage_nesting_material_y_n", "homecage_nesting_material_y_n"],
    "experiment_cohort": ["experiment_cohort"],
    "bred_or_shipped": ["bred_on_site_or_shipped"],
    "age_at_weaning_days": ["age_at_weaning", "age_at_weaning_p"],
    "homecage_hut": ["home_cage_hut_y_n", "homecage_hut_y_n"],
    "sex": ["sex_m_f", "sex"],
    "shipping_age_postnatal_day": ["shipping_age_postnatal_day", "shipping_age", "shipping_arrival_age_p"],
    "homecage_chewing_material": ["home_cage_chewing_material_y_n", "homecage_chewing_material_y_n"],
    "lab_source": ["lab_and_university"],
    "testing_shavings": ["testing_shavings_used_aspen_pine"],
    "testing_light_phase": ["testing_in_dark_or_light_phase"],
    "reinforcer": ["reinforcer_used_bioserv_cheerio"],
    "reinforcer_size_mg": ["reinforcer_size_mg"],
    "colony_bedding": ["colony_bedding"],
    "testing_shavings_manufacturer": ["testing_shavings_manufacturer"],
    "testing_room_lumens": ["lumens_in_testing_room"],
    "experimenter": ["experimenter"],
    "colony_light_cycle": ["colony_light_cycle_e_g_12_12"],
    "testing_arena_dimensions": ["testing_arena_dimensions_lxwxh"],
    "video_recorded": ["video_taken_of_disc_recall_rev_y_n"],
    "video_recorded_discrimination": ["video_taken_of_discrimination_y_n"],
    "video_recorded_recall": ["video_taken_of_recall_y_n"],
    "video_recorded_reversal": ["video_taken_of_reversal_y_n"],
    "colony_lights_on_time": ["colony_lights_on_time_am_pmr", "colony_lights_on_time_am_pm"],
    "testing_ramekin_diameter": ["testing_ramekin_diameter"],
    "video_frame_rate": ["video_frame_rate"],
    "colony_chow_brand": ["colony_chow_brand_fed"],
    "testing_ramekins_sham_baited": ["testing_ramekins_sham_baited_y_n"],
    "video_resolution": ["video_resolution"],
}


def is_formula_value(value: Any) -> bool:
    """Return True for Excel formulas/helper formula objects."""
    return isinstance(value, ArrayFormula) or clean_text(value).startswith("=")


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, ArrayFormula):
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    return str(value).replace("\n", " ").strip()


def clean_label(value: Any) -> str:
    label = re.sub(r"[^a-z0-9]+", "_", clean_text(value).lower()).strip("_")
    return "_".join(part for part in label.split("_") if part)


def is_missing(value: Any) -> bool:
    return clean_text(value).upper() in MISSING_TEXT


def is_no_value(value: Any) -> bool:
    return clean_text(value).upper() in NO_VALUE_TEXT


def maybe_float(value: Any) -> float | None:
    if value is None or value == "" or is_formula_value(value):
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if isinstance(value, float) and math.isnan(value):
            return None
        return float(value)
    text = clean_text(value)
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def maybe_int(value: Any) -> int | None:
    number = maybe_float(value)
    if number is None or not float(number).is_integer():
        return None
    return int(number)


def excel_date(value: Any) -> str:
    if is_formula_value(value):
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    return clean_text(value)


def parse_date(value: Any) -> dt.date | None:
    if is_formula_value(value):
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = clean_text(value)
    if text.upper() in NO_VALUE_TEXT:
        return None
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%m/%d/%Y", "%m/%d/%y"):
        try:
            return dt.datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def parse_time(value: Any) -> dt.time | None:
    if is_formula_value(value):
        return None
    if isinstance(value, dt.datetime):
        return value.time()
    if isinstance(value, dt.time):
        return value
    text = clean_text(value)
    if text.upper() in NO_VALUE_TEXT:
        return None
    for fmt in ("%H:%M:%S", "%H:%M", "%I:%M:%S %p", "%I:%M %p"):
        try:
            return dt.datetime.strptime(text, fmt).time()
        except ValueError:
            pass
    return None


def format_time(value: Any) -> str:
    parsed = parse_time(value)
    if parsed:
        return parsed.strftime("%H:%M")
    return "" if is_formula_value(value) else clean_text(value)


def compute_zt(test_time: Any, lights_on_time: Any) -> str:
    """Compute zeitgeber time in Python from a clock time and lights-on time."""
    test = parse_time(test_time)
    lights_on = parse_time(lights_on_time)
    if not test or not lights_on:
        return ""
    test_minutes = test.hour * 60 + test.minute + test.second / 60
    lights_minutes = lights_on.hour * 60 + lights_on.minute + lights_on.second / 60
    delta = (test_minutes - lights_minutes) % (24 * 60)
    hours = int(delta // 60)
    minutes = int(round(delta % 60))
    if minutes == 60:
        hours = (hours + 1) % 24
        minutes = 0
    return f"ZT {hours:02d}:{minutes:02d}"


def compute_age_days(test_date: Any, dob: Any) -> float | None:
    parsed_test_date = parse_date(test_date)
    parsed_dob = parse_date(dob)
    if not parsed_test_date or not parsed_dob:
        return None
    return float((parsed_test_date - parsed_dob).days)


def sheet_value_pairs(ws, max_row: int = 30, max_col: int = 10) -> dict[str, Any]:
    """Read label/value pairs where labels are followed by values to the right."""
    pairs: dict[str, Any] = {}
    for row in range(1, min(ws.max_row, max_row) + 1):
        for col in range(1, min(ws.max_column, max_col) + 1):
            raw_label = ws.cell(row, col).value
            label = clean_label(raw_label)
            if not label:
                continue
            value = ws.cell(row, col + 1).value if col + 1 <= ws.max_column else None
            if value not in (None, "") and not is_formula_value(value):
                pairs[label] = value
    return pairs


def extract_metadata(wb, source_file: Path | None = None) -> dict[str, Any]:
    if "METADATA" not in wb.sheetnames:
        return {"source_file": str(source_file) if source_file else "", "animal_id": source_file.stem if source_file else "unknown_animal"}
    ws = wb["METADATA"]
    pairs = sheet_value_pairs(ws, max_row=26, max_col=9)
    out: dict[str, Any] = {"source_file": str(source_file) if source_file else ""}
    used_keys: set[str] = set()
    for canonical, keys in METADATA_ALIASES.items():
        for key in keys:
            if key in pairs and key not in used_keys:
                out[canonical] = pairs[key]
                used_keys.add(key)
                break
    out["animal_id_inferred_from_filename"] = False
    if is_missing(out.get("animal_id")):
        out["animal_id"] = source_file.stem if source_file else "unknown_animal"
        out["animal_id_inferred_from_filename"] = True
    out["animal_id"] = clean_text(out["animal_id"])
    for key in ["dob"]:
        if key in out:
            out[key] = excel_date(out[key])
    for key in ["colony_lights_on_time"]:
        if key in out:
            out[key] = format_time(out[key])
    return out


def extract_weight_rows(wb, animal: dict[str, Any], source_file: Path) -> pd.DataFrame:
    if "WEIGHT" not in wb.sheetnames:
        return pd.DataFrame()
    ws = wb["WEIGHT"]
    rows = []
    baseline_weight = maybe_float(ws.cell(10, 2).value)
    dob = animal.get("dob")
    lights_on_time = animal.get("colony_lights_on_time")
    for col in range(2, ws.max_column + 1):
        experimental_day = clean_text(ws.cell(4, col).value)
        stage = clean_text(ws.cell(5, col).value)
        date = excel_date(ws.cell(6, col).value)
        time_value = ws.cell(7, col).value
        time = format_time(time_value)
        zeitgeber = compute_zt(time_value, lights_on_time)
        age = compute_age_days(ws.cell(6, col).value, dob)
        weight_g = maybe_float(ws.cell(10, col).value)
        percent_baseline = None
        if baseline_weight and weight_g is not None:
            percent_baseline = weight_g / baseline_weight
        if not any([date, time, age is not None, weight_g is not None, percent_baseline is not None]):
            continue
        phase = WEIGHT_STAGE_TO_PHASE.get(stage.lower(), stage.lower())
        rows.append(
            {
                "source_file": str(source_file),
                "animal_id": animal.get("animal_id", "unknown_animal"),
                "experimental_day": experimental_day,
                "current_behavioral_phase": phase,
                "stage": stage,
                "date": date,
                "time": time,
                "zeitgeber_time": zeitgeber,
                "age_days": age,
                "weight_g": weight_g,
                "percent_baseline": percent_baseline,
            }
        )
    return pd.DataFrame(rows)


def extract_acclim_rows(wb, animal: dict[str, Any], source_file: Path) -> pd.DataFrame:
    sheet_name = "ACCLIM" if "ACCLIM" in wb.sheetnames else "Acclimation Day 3" if "Acclimation Day 3" in wb.sheetnames else ""
    if not sheet_name:
        return pd.DataFrame()
    ws = wb[sheet_name]
    session_pairs = sheet_value_pairs(ws, max_row=8, max_col=6)
    lights_on_time = animal.get("colony_lights_on_time")
    rows = []
    for row in range(10, ws.max_row + 1):
        round_label = clean_text(ws.cell(row, 1).value)
        values = {
            "time_to_eat_nw_sec": maybe_float(ws.cell(row, 2).value),
            "time_to_eat_ne_sec": maybe_float(ws.cell(row, 3).value),
            "time_to_eat_sw_sec": maybe_float(ws.cell(row, 4).value),
            "time_to_eat_se_sec": maybe_float(ws.cell(row, 5).value),
            "pellets_eaten": clean_text(ws.cell(row, 6).value),
            "notes": clean_text(ws.cell(row, 7).value),
        }
        if not any(v not in (None, "") for v in values.values()):
            continue
        rows.append(
            {
                "source_file": str(source_file),
                "animal_id": animal.get("animal_id", "unknown_animal"),
                "phase": "acclimation",
                "session_date": excel_date(session_pairs.get("date")),
                "start_time": format_time(session_pairs.get("start_time")),
                "end_time": format_time(session_pairs.get("end_time")),
                "start_zeitgeber_time": compute_zt(session_pairs.get("start_time"), lights_on_time),
                "end_zeitgeber_time": compute_zt(session_pairs.get("end_time"), lights_on_time),
                "round": round_label,
                **values,
            }
        )
    return pd.DataFrame(rows)


def metadata_qc(metadata: pd.DataFrame) -> pd.DataFrame:
    required = ["animal_id", "experiment_cohort", "adversity_condition", "sex", "genotype", "lab_source"]
    controlled = {
        "sex": {"M", "F", "UNKNOWN"},
        "adversity_condition": {"CONTROL", "TREATMENT", "N/A", "NA", "NONE", "UNKNOWN"},
    }
    date_fields = ["dob"]
    time_fields = ["colony_lights_on_time"]
    records = []
    for _, row in metadata.iterrows():
        animal_id = clean_text(row.get("animal_id")) or "unknown_animal"
        source_file = clean_text(row.get("source_file"))
        for field in required:
            value = clean_text(row.get(field))
            if is_missing(value) and field != "adversity_condition":
                records.append(
                    qc_record(
                        animal_id,
                        source_file,
                        field,
                        "missing_required_metadata",
                        "",
                        f"Missing required metadata field '{field}'. Fill this field in the METADATA sheet.",
                    )
                )
        if bool(row.get("animal_id_inferred_from_filename", False)):
            records.append(
                qc_record(
                    animal_id,
                    source_file,
                    "animal_id",
                    "animal_id_inferred_from_filename",
                    animal_id,
                    "Animal ID was blank, so the pipeline used the workbook filename. Enter the animal ID in METADATA if this is not intended.",
                )
            )
        for field, allowed in controlled.items():
            value = clean_text(row.get(field))
            if value and value.upper() not in allowed:
                records.append(
                    qc_record(
                        animal_id,
                        source_file,
                        field,
                        "nonstandard_controlled_value",
                        value,
                        f"Field '{field}' has value '{value}', which is outside the expected values: {', '.join(sorted(allowed))}.",
                    )
                )
        for field in date_fields:
            value = row.get(field)
            if not is_no_value(value) and parse_date(value) is None:
                records.append(
                    qc_record(
                        animal_id,
                        source_file,
                        field,
                        "invalid_date_format",
                        clean_text(value),
                        f"Field '{field}' should be a real spreadsheet date or YYYY-MM-DD date.",
                    )
                )
        for field in time_fields:
            value = row.get(field)
            if not is_no_value(value) and parse_time(value) is None:
                records.append(
                    qc_record(
                        animal_id,
                        source_file,
                        field,
                        "invalid_time_format",
                        clean_text(value),
                        f"Field '{field}' should be a real spreadsheet time or HH:MM time.",
                    )
                )
    if "animal_id" in metadata.columns:
        counts = metadata["animal_id"].map(clean_text).value_counts()
        for animal_id, count in counts[counts > 1].items():
            records.append(
                qc_record(
                    animal_id,
                    "",
                    "animal_id",
                    "duplicate_animal_id",
                    str(count),
                    f"Animal ID '{animal_id}' appears in {count} workbooks. Confirm this is intentional before combining data.",
                )
            )
    return pd.DataFrame(records, columns=["animal_id", "source_file", "field", "qc_flag", "value", "message"])


def qc_record(animal_id: str, source_file: str, field: str, qc_flag: str, value: str, message: str) -> dict[str, str]:
    return {
        "animal_id": animal_id,
        "source_file": source_file,
        "field": field,
        "qc_flag": qc_flag,
        "value": value,
        "message": message,
    }


def submission_qc(metadata: pd.DataFrame, trials: pd.DataFrame, weights: pd.DataFrame) -> pd.DataFrame:
    """QC checks that require more than one parsed table."""
    records = []
    if "animal_id" not in metadata.columns:
        return pd.DataFrame(columns=["animal_id", "source_file", "field", "qc_flag", "value", "message"])

    metadata_animals = set(metadata["animal_id"].map(clean_text))
    trial_animals = set(trials["animal_id"].map(clean_text)) if "animal_id" in trials.columns else set()
    weight_animals = set(weights["animal_id"].map(clean_text)) if "animal_id" in weights.columns else set()
    source_by_animal = metadata.set_index(metadata["animal_id"].map(clean_text))["source_file"].map(clean_text).to_dict() if "source_file" in metadata.columns else {}

    for animal_id in sorted(metadata_animals - trial_animals):
        records.append(
            qc_record(
                animal_id,
                source_by_animal.get(animal_id, ""),
                "trial_level",
                "no_behavior_trials_found",
                "",
                "No behavior trial rows were found for this animal. Fill Odor Chosen and/or Latency in DISC, RECALL, or REV if behavior data should be present.",
            )
        )
    for animal_id in sorted(metadata_animals - weight_animals):
        records.append(
            qc_record(
                animal_id,
                source_by_animal.get(animal_id, ""),
                "weight_long",
                "no_weight_rows_found",
                "",
                "No weight rows were found for this animal. Fill Date, Time, or Weight in the WEIGHT sheet if weight data should be present.",
            )
        )
    return pd.DataFrame(records, columns=["animal_id", "source_file", "field", "qc_flag", "value", "message"])


def odor_labels(ws) -> dict[str, str]:
    labels = {}
    for col in range(2, 6):
        text = clean_text(ws.cell(7, col).value)
        if ":" in text:
            code, name = text.split(":", 1)
            labels[clean_text(code)] = clean_text(name)
    return labels


def rewarded_odor(ws, phase: str) -> str:
    explicit = clean_text(ws.cell(6, 5).value)
    labels = odor_labels(ws)
    if explicit:
        for code, label in labels.items():
            if clean_text(label).lower() == explicit.lower():
                return code
        if re.fullmatch(r"O[1-4]", explicit, re.I):
            return explicit.upper()
        if re.fullmatch(r"[A-D]", explicit, re.I):
            return explicit.upper()
    if any(code in labels for code in ["A", "B", "C", "D"]):
        return "B" if phase == "reversal" else "A"
    return "O2" if phase == "reversal" else "O1"


def normalized_choice(value: Any) -> str:
    text = clean_text(value)
    if not text:
        return ""
    if text.lower() in {"omit", "omission", "miss", "no choice"}:
        return "Omit"
    upper = text.upper()
    if upper in {"O1", "O2", "O3", "O4"}:
        return upper
    if upper in {"A", "B", "C", "D"}:
        return upper
    return text


def quadrant_from_order(order: str, odor: str) -> str:
    if odor == "Omit" or not odor:
        return ""
    normalized_order = ", ".join(part.strip() for part in clean_text(order).split(","))
    mapping = ODOR_TO_QUADRANT.get(normalized_order)
    if mapping:
        return mapping.get(odor, "")
    parts = [part.strip() for part in clean_text(order).split(",")]
    if len(parts) == 4 and odor in parts:
        return ["NW", "NE", "SW", "SE"][parts.index(odor)]
    return ""


def odor_from_quadrant(order: str, quadrant: str) -> str:
    if not quadrant:
        return ""
    normalized_order = clean_text(order).replace(" ", "")
    mapping = ODOR_TO_QUADRANT.get(normalized_order)
    if not mapping:
        normalized_order = ", ".join(part.strip() for part in clean_text(order).split(","))
        mapping = ODOR_TO_QUADRANT.get(normalized_order, {})
    reverse = {quad: odor for odor, quad in mapping.items()}
    return reverse.get(quadrant, "")


def parse_entry_sequence(entry_sequence: Any, order: str) -> tuple[list[str], list[str]]:
    text = clean_text(entry_sequence)
    quadrants: list[str] = []
    for char in text:
        if char in QUADRANT_NUMBER:
            quadrants.append(QUADRANT_NUMBER[char])
        elif char.upper() in {"N", "E", "S", "W"}:
            continue
    odors = [odor_from_quadrant(order, quadrant) for quadrant in quadrants]
    return quadrants, odors


def reversal_error_type(choice: str, previous_rewarded: str, new_rewarded: str, seen_correct: bool = False) -> str:
    if choice in {"", "Omit", new_rewarded}:
        return ""
    if choice == previous_rewarded:
        return "Regressive" if seen_correct else "Perseverative"
    irrelevant = "C" if choice in {"A", "B", "C", "D"} else "O3"
    novel = "D" if choice in {"A", "B", "C", "D"} else "O4"
    if choice == irrelevant:
        return "Irrelevant"
    if choice == novel:
        return "Novel"
    return "Other"


def parse_trial_sheet(wb, sheet_name: str, animal: dict[str, Any], source_file: Path) -> list[dict[str, Any]]:
    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]
    phase = TRIAL_SHEETS[sheet_name]
    rewarded = rewarded_odor(ws, phase)
    previous_rewarded = "A" if rewarded in {"A", "B", "C", "D"} else "O1"
    labels = odor_labels(ws)
    session_pairs = sheet_value_pairs(ws, max_row=8, max_col=6)
    lights_on_time = animal.get("colony_lights_on_time")
    headers = {clean_label(ws.cell(9, col).value): col for col in range(1, ws.max_column + 1) if clean_label(ws.cell(9, col).value)}
    choice_col = headers.get("odor_chosen", 4)
    latency_col = headers.get("latency_s", 6)
    entry_sequence_col = headers.get("entry_sequence", 7)
    notes_col = headers.get("notes", 10 if sheet_name in {"REV", "REVERSAL"} else 9)
    rows = []
    seen_reversal_correct = False

    for row in range(10, ws.max_row + 1):
        trial = maybe_int(ws.cell(row, 1).value)
        if trial is None:
            continue
        choice = normalized_choice(ws.cell(row, choice_col).value)
        if not choice:
            # Blank column D is the agreed way to skip a mistaken row.
            continue
        latency = maybe_float(ws.cell(row, latency_col).value)
        entry_sequence = clean_text(ws.cell(row, entry_sequence_col).value)
        order = clean_text(ws.cell(row, 2).value)
        entry_quadrants, entry_odors = parse_entry_sequence(entry_sequence, order)
        entries = len(entry_quadrants) if entry_quadrants else maybe_float(ws.cell(row, 7).value)

        has_data = any([choice, latency is not None, entries is not None, entry_sequence])
        if not has_data:
            continue

        is_omission = choice == "Omit"
        is_correct = choice == rewarded if choice and not is_omission else False
        is_error = bool(choice and not is_omission and not is_correct)
        error_type = ""
        if phase == "reversal":
            error_type = reversal_error_type(choice, previous_rewarded=previous_rewarded, new_rewarded=rewarded, seen_correct=seen_reversal_correct)
            if is_correct:
                seen_reversal_correct = True
        elif is_omission:
            error_type = "Omission"
        elif is_error:
            error_type = f"{choice}_error" if choice in {"A", "B", "C", "D", "O1", "O2", "O3", "O4"} else "Other"

        rows.append(
            {
                "source_file": str(source_file),
                "animal_id": animal.get("animal_id", "unknown_animal"),
                "lab_source": clean_text(animal.get("lab_source")),
                "experiment_cohort": clean_text(animal.get("experiment_cohort")),
                "group": clean_text(animal.get("adversity_condition")),
                "sex": clean_text(animal.get("sex")),
                "genotype": clean_text(animal.get("genotype")),
                "strain": clean_text(animal.get("strain")),
                "phase": phase,
                "sheet_name": sheet_name,
                "session_date": excel_date(session_pairs.get("date")),
                "session_start_time": format_time(session_pairs.get("start_time")),
                "session_end_time": format_time(session_pairs.get("end_time")),
                "session_start_zeitgeber_time": compute_zt(session_pairs.get("start_time"), lights_on_time),
                "session_end_zeitgeber_time": compute_zt(session_pairs.get("end_time"), lights_on_time),
                "trial": trial,
                "excel_row": row,
                "bowl_order": order,
                "odor_chosen": choice,
                "odor_chosen_label": labels.get(choice, choice),
                "quadrant_chosen": quadrant_from_order(order, choice),
                "rewarded_odor": rewarded,
                "rewarded_label": labels.get(rewarded, rewarded),
                "correct": bool(is_correct),
                "error": bool(is_error),
                "omission": bool(is_omission),
                "error_type": error_type,
                "latency_sec": latency,
                "entries": entries,
                "entry_sequence": entry_sequence,
                "entry_quadrant_sequence": ",".join(entry_quadrants),
                "entry_odor_sequence": ",".join(entry_odors),
                "entry_A_count": entry_odors.count("A"),
                "entry_B_count": entry_odors.count("B"),
                "entry_C_count": entry_odors.count("C"),
                "entry_D_count": entry_odors.count("D"),
                "entry_O1_count": entry_odors.count("O1"),
                "entry_O2_count": entry_odors.count("O2"),
                "entry_O3_count": entry_odors.count("O3"),
                "entry_O4_count": entry_odors.count("O4"),
                "notes": clean_text(ws.cell(row, notes_col).value),
            }
        )
    return rows


def parse_tabs_workbook(path: Path) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    wb = load_workbook(path, data_only=False)
    animal = extract_metadata(wb, source_file=path)
    weight_df = extract_weight_rows(wb, animal, path)
    acclim_df = extract_acclim_rows(wb, animal, path)
    trials = []
    for sheet_name in TRIAL_SHEETS:
        trials.extend(parse_trial_sheet(wb, sheet_name, animal, path))
    trial_df = pd.DataFrame(trials)
    if trial_df.empty:
        return trial_df, pd.DataFrame([animal]), weight_df, acclim_df
    trial_df = trial_df.sort_values(["animal_id", "phase", "trial", "excel_row"]).reset_index(drop=True)
    trial_df = add_sequence_metrics(trial_df)
    return trial_df, pd.DataFrame([animal]), weight_df, acclim_df


def add_sequence_metrics(trials: pd.DataFrame, criterion_correct: int = 8, window_trials: int = 10) -> pd.DataFrame:
    trials = trials.copy()
    trials["correct_running_count"] = 0
    trials["non_omission_count"] = 0
    trials["correct_in_last_10_non_omission"] = np.nan
    trials["reversal_first_correct_seen"] = False
    out = []
    for (_, phase), sub in trials.groupby(["animal_id", "phase"], sort=False):
        run = 0
        seen = False
        non_omission_correct: list[bool] = []
        for _, row in sub.iterrows():
            if bool(row["correct"]):
                run += 1
                if phase == "reversal":
                    seen = True
            else:
                run = 0
            if not bool(row["omission"]):
                non_omission_correct.append(bool(row["correct"]))
            row["correct_running_count"] = run
            row["non_omission_count"] = len(non_omission_correct)
            if len(non_omission_correct) >= window_trials:
                row["correct_in_last_10_non_omission"] = int(sum(non_omission_correct[-window_trials:]))
            row["reversal_first_correct_seen"] = seen
            out.append(row)
    return pd.DataFrame(out)


def trials_to_criterion(sub: pd.DataFrame, criterion_correct: int = 8, window_trials: int = 10) -> int | None:
    if sub.empty:
        return None
    if "correct_in_last_10_non_omission" not in sub.columns:
        sub = add_sequence_metrics(sub, criterion_correct=criterion_correct, window_trials=window_trials)
    hit = sub.loc[sub["correct_in_last_10_non_omission"] >= criterion_correct, "trial"]
    if len(hit):
        return int(hit.iloc[0])
    non_omission = sub.loc[~sub["omission"].astype(bool), "trial"]
    if len(non_omission):
        return int(non_omission.max())
    return int(sub["trial"].max())


def summarize_sessions(trials: pd.DataFrame) -> pd.DataFrame:
    if trials.empty:
        return pd.DataFrame()
    if "correct_running_count" not in trials.columns:
        trials = add_sequence_metrics(trials)
    records = []
    group_cols = ["animal_id", "phase"]
    for (animal_id, phase), sub in trials.groupby(group_cols, sort=False):
        record = {
            "animal_id": animal_id,
            "phase": phase,
            "lab_source": first_nonempty(sub["lab_source"]),
            "experiment_cohort": first_nonempty(sub["experiment_cohort"]),
            "group": first_nonempty(sub["group"]),
            "sex": first_nonempty(sub["sex"]),
            "genotype": first_nonempty(sub["genotype"]),
            "session_date": first_nonempty(sub["session_date"]),
            "n_trials": int(len(sub)),
            "trials_to_criterion_8of10_non_omission": trials_to_criterion(sub),
            "total_errors": int(sub["error"].sum()),
            "omissions": int(sub["omission"].sum()),
            "mean_latency_sec": sub["latency_sec"].dropna().mean(),
            "median_latency_sec": sub["latency_sec"].dropna().median(),
            "mean_latency_correct_sec": sub.loc[sub["correct"], "latency_sec"].dropna().mean(),
            "mean_latency_incorrect_sec": sub.loc[sub["error"], "latency_sec"].dropna().mean(),
            "mean_entries": sub["entries"].dropna().mean(),
            "mean_entries_correct": sub.loc[sub["correct"], "entries"].dropna().mean(),
            "mean_entries_incorrect": sub.loc[sub["error"], "entries"].dropna().mean(),
            "total_entries": sub["entries"].dropna().sum(),
        }
        odor_codes = ["A", "B", "C", "D"] if sub["odor_chosen"].isin(["A", "B", "C", "D"]).any() else ["O1", "O2", "O3", "O4"]
        for code in odor_codes:
            record[f"{code}_errors"] = int(((sub["odor_chosen"] == code) & sub["error"]).sum())
            record[f"{code}_choices"] = int((sub["odor_chosen"] == code).sum())
            record[f"{code}_choice_percent_trials"] = float((sub["odor_chosen"] == code).sum() / len(sub) * 100) if len(sub) else np.nan
            entry_col = f"entry_{code}_count"
            if entry_col in sub:
                total_entry_events = sub[[f"entry_{c}_count" for c in odor_codes if f"entry_{c}_count" in sub]].sum().sum()
                record[f"{code}_entries"] = int(sub[entry_col].sum())
                record[f"{code}_entry_percent_total"] = float(sub[entry_col].sum() / total_entry_events * 100) if total_entry_events else np.nan
        for error_type in ["Perseverative", "Regressive", "Irrelevant", "Novel", "Omission", "Other"]:
            record[f"{clean_label(error_type)}_errors"] = int((sub["error_type"] == error_type).sum())
        records.append(record)
    return pd.DataFrame(records)


def summarize_animals(trials: pd.DataFrame, metadata: pd.DataFrame) -> pd.DataFrame:
    sessions = summarize_sessions(trials)
    metadata_base = metadata.copy()
    if "group" not in metadata_base.columns and "adversity_condition" in metadata_base.columns:
        metadata_base["group"] = metadata_base["adversity_condition"]
    if "animal_id" in metadata_base.columns:
        metadata_base = metadata_base.drop_duplicates("animal_id", keep="first")
    if sessions.empty:
        return metadata_base
    base_cols = ["animal_id", "lab_source", "experiment_cohort", "group", "sex", "genotype"]
    session_animals = sessions.groupby("animal_id", as_index=False).agg({col: first_nonempty for col in base_cols[1:]})
    if "animal_id" in metadata_base.columns and not metadata_base.empty:
        animals = metadata_base.merge(session_animals, on="animal_id", how="outer", suffixes=("", "_from_trials"))
        for col in base_cols[1:]:
            trial_col = f"{col}_from_trials"
            if trial_col in animals.columns:
                if col in animals.columns:
                    animals[col] = animals[col].where(animals[col].map(clean_text).astype(bool), animals[trial_col])
                else:
                    animals[col] = animals[trial_col]
                animals = animals.drop(columns=[trial_col])
    else:
        animals = session_animals
    metric_cols = [
        "n_trials",
        "trials_to_criterion_8of10_non_omission",
        "total_errors",
        "omissions",
        "mean_latency_sec",
        "median_latency_sec",
        "mean_latency_correct_sec",
        "mean_latency_incorrect_sec",
        "mean_entries",
        "mean_entries_correct",
        "mean_entries_incorrect",
        "total_entries",
        "A_errors",
        "B_errors",
        "C_errors",
        "D_errors",
        "A_choices",
        "B_choices",
        "C_choices",
        "D_choices",
        "A_choice_percent_trials",
        "B_choice_percent_trials",
        "C_choice_percent_trials",
        "D_choice_percent_trials",
        "A_entries",
        "B_entries",
        "C_entries",
        "D_entries",
        "A_entry_percent_total",
        "B_entry_percent_total",
        "C_entry_percent_total",
        "D_entry_percent_total",
        "O1_errors",
        "O2_errors",
        "O3_errors",
        "O4_errors",
        "perseverative_errors",
        "regressive_errors",
        "irrelevant_errors",
        "novel_errors",
    ]
    new_columns = {}
    for phase, prefix in [("discrimination", "disc"), ("recall", "recall"), ("reversal", "rev")]:
        sub = sessions[sessions["phase"] == phase].set_index("animal_id")
        for col in metric_cols:
            if col in sub.columns:
                new_columns[f"{prefix}_{col}"] = animals["animal_id"].map(sub[col])
    if new_columns:
        animals = pd.concat([animals, pd.DataFrame(new_columns, index=animals.index)], axis=1)
    summary_columns = {
        "TTC_in_discrimination": animals.get("disc_trials_to_criterion_8of10_non_omission"),
        "Total_errors_in_discrimination": animals.get("disc_total_errors"),
        "TTC_in_reversal": animals.get("rev_trials_to_criterion_8of10_non_omission"),
        "Total_errors_in_reversal": animals.get("rev_total_errors"),
        "Perseverative_errors_in_reversal": animals.get("rev_perseverative_errors"),
        "Regressive_errors_in_reversal": animals.get("rev_regressive_errors"),
        "O1_errors": animals.get("rev_O1_errors"),
    }
    if "rev_A_errors" in animals:
        summary_columns["O1_errors"] = animals.get("rev_A_errors")
        summary_columns["O3_errors"] = animals.get("rev_C_errors")
        summary_columns["O4_errors"] = animals.get("rev_D_errors")
    else:
        summary_columns["O3_errors"] = animals.get("rev_O3_errors")
        summary_columns["O4_errors"] = animals.get("rev_O4_errors")
    animals = pd.concat([animals, pd.DataFrame(summary_columns, index=animals.index)], axis=1)
    return animals


def first_nonempty(values: Any) -> str:
    if isinstance(values, pd.Series):
        iterator = values
    else:
        iterator = [values]
    for value in iterator:
        text = clean_text(value)
        if text:
            return text
    return ""


def pca_table(animals: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    features = [
        "TTC_in_discrimination",
        "Total_errors_in_discrimination",
        "TTC_in_reversal",
        "Total_errors_in_reversal",
        "Perseverative_errors_in_reversal",
        "Regressive_errors_in_reversal",
        "O1_errors",
        "O3_errors",
        "O4_errors",
    ]
    use = [col for col in features if col in animals.columns and animals[col].notna().any()]
    if len(use) < 2 or len(animals) < 2:
        return pd.DataFrame(), pd.DataFrame(), pd.DataFrame()
    x = animals[use].apply(pd.to_numeric, errors="coerce")
    x = x.fillna(x.mean())
    std = x.std(ddof=0).replace(0, 1)
    z = (x - x.mean()) / std
    u, s, vt = np.linalg.svd(z.to_numpy(), full_matrices=False)
    ncomp = min(3, vt.shape[0])
    scores = pd.DataFrame(u[:, :ncomp] * s[:ncomp], columns=[f"PC{i}" for i in range(1, ncomp + 1)])
    scores.insert(0, "animal_id", animals["animal_id"].values)
    loadings = pd.DataFrame(vt[:ncomp].T, index=use, columns=[f"PC{i}" for i in range(1, ncomp + 1)]).reset_index()
    loadings = loadings.rename(columns={"index": "feature"})
    variance = pd.DataFrame(
        {
            "component": [f"PC{i}" for i in range(1, ncomp + 1)],
            "explained_variance_ratio": (s[:ncomp] ** 2) / (s**2).sum(),
        }
    )
    return scores, loadings, variance


def autosize_excel(path: Path) -> None:
    wb = load_workbook(path)
    for ws in wb.worksheets:
        for col in range(1, min(ws.max_column, 60) + 1):
            letter = get_column_letter(col)
            max_len = 8
            for cell in ws[letter][: min(ws.max_row, 200)]:
                max_len = max(max_len, min(40, len(clean_text(cell.value))))
            ws.column_dimensions[letter].width = max_len + 2
        ws.freeze_panes = "A2"
    wb.save(path)
